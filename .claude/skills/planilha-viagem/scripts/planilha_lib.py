"""
Biblioteca compartilhada por gerar-planilha.py e comparar-viagens.py.

Transforma um trip.json em abas de planilha com fórmulas vivas. Nenhum total é
escrito como número: tudo que é derivado é fórmula, para que mudar um preço ou
trocar a opção de um item refaça o Resumo, os Cortes e a Comparação sozinho.

Cada item de custo tem três opções — econômico (downgrade), plano e upgrade —
com preço, descrição e link de compra. A coluna "Opção" escolhe qual entra no
total. É assim que o usuário monta o próprio cenário sem fórmula nenhuma.

Convenções (iguais às de modelo financeiro):
  azul   = número de entrada (edite à vontade)
  preto  = fórmula (não edite)
  verde  = fórmula que lê outra aba
  fundo amarelo = premissa-chave (teto, pessoas, reserva, câmbio)
"""
from __future__ import annotations

import json
from datetime import date, datetime
from pathlib import Path

from openpyxl.chart import BarChart, Reference
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

FONTE = "Arial"
AZUL = "0000FF"
PRETO = "000000"
VERDE = "008000"
CINZA = "666666"
LINK = "0563C1"
AMARELO = PatternFill("solid", fgColor="FFFF00")
CABECALHO = PatternFill("solid", fgColor="1F2A44")
SUBTOTAL = PatternFill("solid", fgColor="EEF1F7")
INTOCAVEL = PatternFill("solid", fgColor="FDECEC")
ESCOLHA = PatternFill("solid", fgColor="FFF6D5")

FMT_EUR = '"€ "#,##0.00;-"€ "#,##0.00;"-"'
FMT_EUR0 = '"€ "#,##0;-"€ "#,##0;"-"'
FMT_NUM = "#,##0.00"
FMT_PCT = "0.0%"
FMT_DATA = "DD/MM/YYYY"
FMT_TAXA = "0.000000"

OPCOES = ("econômico", "plano", "upgrade")

CATEGORIAS_PADRAO = [
    ("voos", "Voos"), ("hospedagem", "Hospedagem"), ("taxa-turismo", "Taxa de turismo"),
    ("alimentacao", "Alimentação"), ("transporte-local", "Transporte local"),
    ("atracoes", "Atrações"), ("documentacao", "Documentação"),
    ("conectividade", "Conectividade"), ("compras", "Compras"), ("extras", "Extras"),
]
TIPO_PARA_CATEGORIA = {
    "atracao": "atracoes", "evento": "atracoes", "refeicao": "alimentacao",
    "deslocamento": "transporte-local", "compras": "compras", "livre": "extras",
}
SEMANA = ["segunda", "terça", "quarta", "quinta", "sexta", "sábado", "domingo"]

# Colunas da aba Itens (letra → papel). Mudar aqui muda em todas as fórmulas.
COL = dict(
    id="A", cat="B", item="C", quando="D", qtd="E", unid="F", moeda="G", opcao="H",
    p_eco="I", p_plano="J", p_up="K", p_esc="L", mult="M", cambio="N",
    t_eco="O", t_plano="P", t_up="Q", t_esc="R", intoc="S", dia="T",
    link="U", d_eco="V", d_up="W", fonte="X",
)


# ── utilidades ─────────────────────────────────────────────────────────────

def ler_trip(raiz: Path, slug: str) -> dict:
    return json.loads((raiz / "destinos" / slug / "trip.json").read_text("utf-8"))


def iso(d: str | None) -> date | None:
    if not d:
        return None
    return datetime.strptime(str(d)[:10], "%Y-%m-%d").date()


def q(nome: str) -> str:
    """Nome de aba entre aspas para usar em fórmula."""
    return "'" + nome.replace("'", "''") + "'"


def fonte(cor=PRETO, negrito=False, italico=False, tamanho=10):
    return Font(name=FONTE, color=cor, bold=negrito, italic=italico, size=tamanho)


def entrada(cel, valor, fmt=None, chave=False):
    cel.value = valor
    cel.font = fonte(AZUL)
    if fmt:
        cel.number_format = fmt
    if chave:
        cel.fill = AMARELO


def formula(cel, f, fmt=None, cor=PRETO, negrito=False):
    cel.value = f
    cel.font = fonte(cor, negrito)
    if fmt:
        cel.number_format = fmt


def rotulo(cel, texto, negrito=False, cor=PRETO, italico=False):
    cel.value = texto
    cel.font = fonte(cor, negrito, italico)


def link(cel, url: str | None, texto: str | None = None):
    if not url:
        return
    cel.value = texto or url
    cel.hyperlink = url
    cel.font = Font(name=FONTE, color=LINK, underline="single", size=10)


def cabecalho(ws, linha, titulos, larguras=None):
    for i, t in enumerate(titulos, start=1):
        c = ws.cell(row=linha, column=i, value=t)
        c.font = Font(name=FONTE, bold=True, color="FFFFFF", size=10)
        c.fill = CABECALHO
        c.alignment = Alignment(vertical="center", wrap_text=True)
    if larguras:
        for i, w in enumerate(larguras, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[linha].height = 30


def titulo(ws, texto, sub=None):
    ws["A1"].value = texto
    ws["A1"].font = fonte(negrito=True, tamanho=14)
    if sub:
        ws["A2"].value = sub
        ws["A2"].font = fonte(CINZA, italico=True)


def curto(slug: str, n: int = 20) -> str:
    return slug.replace("-", "_")[:n]


def pessoas_de(v: dict) -> int:
    return max(1, len(v.get("viajantes") or [1]))


# ── modelo: itens de custo com três opções ─────────────────────────────────

def noites_hospedagem(h: dict) -> int:
    if h.get("noites"):
        return int(h["noites"])
    a, b = iso(h.get("checkin")), iso(h.get("checkout"))
    return max(1, (b - a).days) if a and b else 1


def _opcoes(c: dict, divisor: float = 1.0) -> dict:
    """Lê custo.opcoes.{economico,upgrade} = {valor, descricao, link}."""
    o = c.get("opcoes") or {}
    out = {}
    for k in ("economico", "upgrade"):
        x = o.get(k) or {}
        val = x.get("valor")
        if val is None and c.get(k) is not None:  # atalho: custo.economico / custo.upgrade
            val = c.get(k)
        out[k] = dict(
            valor=(round(val / divisor, 2) if (val is not None and divisor != 1) else val),
            descricao=x.get("descricao") or "",
            link=x.get("link"),
        )
    return out


def _item(id, categoria, item, quando, qtd, unidade, c, base, fonte_txt="", link_url=None, dia=None, divisor=1.0):
    ops = _opcoes(c, divisor)
    return dict(
        id=id, categoria=categoria, item=item, quando=quando, qtd=qtd, unidade=unidade,
        moeda=c.get("moeda", base), plano=(round(c["valor"] / divisor, 2) if divisor != 1 else c["valor"]),
        economico=ops["economico"]["valor"], upgrade=ops["upgrade"]["valor"],
        d_eco=ops["economico"]["descricao"], d_up=ops["upgrade"]["descricao"],
        link=c.get("link") or link_url or ops["economico"]["link"] or ops["upgrade"]["link"],
        opcao=c.get("opcao", "plano"), fonte=fonte_txt, dia=dia,
    )


def extrair_itens(v: dict) -> list[dict]:
    base = v.get("moedaBase") or v.get("casa", {}).get("moeda", "EUR")
    itens: list[dict] = []

    def custo(c):
        return c if (c and c.get("valor") is not None) else None

    for f in v.get("voos", []):
        c = custo(f.get("custo"))
        if c:
            itens.append(_item(f"voo:{f['id']}", "voos",
                               f"Voo {f.get('de', '')} → {f.get('para', '')} ({f.get('cia') or 'cia a definir'})",
                               iso(f.get("partida")), 1, c.get("por", "pessoa"), c, base,
                               f"status: {f.get('status', 'cotado')}", f.get("link")))

    for h in v.get("hospedagens", []):
        c = custo(h.get("custo"))
        if not c:
            continue
        n = noites_hospedagem(h)
        por = c.get("por", "total")
        if por == "total":
            itens.append(_item(f"hosp:{h['id']}", "hospedagem", h.get("nome", "Hospedagem"), iso(h.get("checkin")),
                               n, "noite", c, base, h.get("porque") or "", h.get("link"), divisor=n))
        else:
            itens.append(_item(f"hosp:{h['id']}", "hospedagem", h.get("nome", "Hospedagem"), iso(h.get("checkin")),
                               n if por == "noite" else 1, "noite" if por == "noite" else "pessoa", c, base,
                               h.get("porque") or "", h.get("link")))

    dest = {d["id"]: d for d in v.get("destinos", [])}
    for h in v.get("hospedagens", []):
        d = dest.get(h.get("destinoId"))
        t = d.get("taxaTurismo") if d else None
        if not t or t.get("valor") is None:
            continue
        itens.append(_item(f"taxa:{h['id']}", "taxa-turismo", f"Taxa de turismo — {d['nome']}", iso(h.get("checkin")),
                           noites_hospedagem(h), "pessoa-noite" if t.get("por", "pessoa-noite") == "pessoa-noite" else "noite",
                           t, base, "cobrada no check-in, por pessoa por noite"))

    for t in v.get("transportes", []):
        c = custo(t.get("custo"))
        if c:
            itens.append(_item(f"transp:{t['id']}", "transporte-local", t.get("nome", "Transporte"), None, 1,
                               c.get("por", "pessoa"), c, base, t.get("observacao") or "", t.get("link")))

    # Bloco de deslocamento com custo só entra quando a viagem não tem
    # transportes[]: se tem, o passe/bilhete já está lá e o bloco é informativo.
    tem_transportes = bool(v.get("transportes"))
    for i, dia in enumerate(v.get("dias", []), start=1):
        for b in dia.get("blocos", []):
            c = custo(b.get("custo"))
            if c and not (tem_transportes and b.get("tipo") == "deslocamento"):
                cat = TIPO_PARA_CATEGORIA.get(b.get("tipo"), "atracoes")
                itens.append(_item(f"dia{i}:{b.get('id', '')}", cat, b.get("titulo", "Bloco"), iso(dia.get("data")), 1,
                                   c.get("por", "pessoa"), c, base, (b.get("notas") or "")[:140], b.get("link"), dia=i))
    return itens


def categorias_do_trip(v: dict):
    cats = []
    nomes = dict(CATEGORIAS_PADRAO)
    for c in v.get("orcamento", {}).get("categorias", []):
        prev = c.get("previsto") or 0
        eco = c.get("economico")
        fator = 1.0 if (c.get("intocavel") or not prev or eco is None) else round(eco / prev, 4)
        cats.append(dict(id=c["id"], nome=c.get("nome") or nomes.get(c["id"], c["id"]), previsto=prev,
                         economico=eco if eco is not None else prev, fator=fator,
                         intocavel=bool(c.get("intocavel")), observacao=c.get("observacao") or "", link=c.get("link")))
    return cats, {c["id"] for c in cats}


def itens_com_restante(v: dict, pessoas: int):
    """Itens derivados + uma linha de estimativa por categoria para fechar com o previsto do trip.json."""
    itens = extrair_itens(v)
    cats, ids = categorias_do_trip(v)
    nomes = dict(CATEGORIAS_PADRAO)
    for it in itens:
        if it["categoria"] not in ids:
            cats.append(dict(id=it["categoria"], nome=nomes.get(it["categoria"], it["categoria"]), previsto=0,
                             economico=0, fator=1.0, intocavel=False, observacao="", link=None))
            ids.add(it["categoria"])
    base = v.get("moedaBase") or v.get("casa", {}).get("moeda", "EUR")
    taxas = {base: 1.0}
    for c in v.get("orcamento", {}).get("cambioReferencia", []):
        if c.get("de") and c.get("taxa"):
            taxas[c["de"]] = float(c["taxa"])
    margem = float(v.get("orcamento", {}).get("margemCartao") or 0)

    def eur(it, preco):
        mult = it["qtd"] * (pessoas if it["unidade"] in ("pessoa", "pessoa-noite") else 1)
        return mult * preco * taxas.get(it["moeda"], 1.0) * (1 + (0 if it["moeda"] == base else margem))

    for c in cats:
        soma = sum(eur(it, it["plano"]) for it in itens if it["categoria"] == c["id"])
        resto = round(c["previsto"] - soma, 2)
        if abs(resto) >= 0.5 or soma == 0:
            itens.append(dict(
                id=f"est:{c['id']}", categoria=c["id"],
                item=(f"Estimativa restante — {c['nome']}" if resto >= 0
                      else f"Ajuste — {c['nome']} (itens somam mais que o previsto)"),
                quando=None, qtd=1, unidade="total", moeda=base, plano=resto, economico=None, upgrade=None,
                d_eco="", d_up="", link=c.get("link"), opcao="plano", fonte=c["observacao"], dia=None, estimativa=True,
            ))
    return itens, cats


# ── abas ───────────────────────────────────────────────────────────────────

class Abas:
    def __init__(self, params="Parâmetros", itens="Itens", real="Real", resumo="Resumo"):
        self.params, self.itens, self.real, self.resumo = params, itens, real, resumo
        self.P, self.I, self.R = q(params), q(itens), q(real)
        self.linhas_itens: dict[str, int] = {}
        self.fim_itens = 2
        self.cambio = (0, 0)
        self.cats = (0, 0)
        self.cat_ids: list[str] = []


def ref_param(abas: Abas, celula: str) -> str:
    return f"{abas.P}!${celula[0]}${celula[1:]}"


def aba_parametros(ws, v: dict, abas: Abas, pessoas: int):
    o = v.get("orcamento", {})
    base = v.get("moedaBase") or v.get("casa", {}).get("moeda", "EUR")
    titulo(ws, v.get("titulo", v["slug"]), "Premissas da viagem. Azul = entrada; amarelo = o que mais mexe no total.")
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 14
    ws.column_dimensions["D"].width = 12
    ws.column_dimensions["E"].width = 70

    linhas = [
        ("Slug", v["slug"], None, None, False),
        ("Início", iso(v["periodo"]["inicio"]), FMT_DATA, None, False),
        ("Fim", iso(v["periodo"]["fim"]), FMT_DATA, None, False),
        ("Noites", "=B5-B4", "0", "fórmula", False),
        ("Dias de viagem", "=B5-B4+1", "0", "fórmula", False),
        ("Pessoas", pessoas, "0", None, True),
        ("Moeda de casa", base, None, None, False),
        ("Teto informado", (o.get("teto") or {}).get("valor"), FMT_EUR0, None, True),
        ("Teto é por", o.get("tetoPor", "total"), None, "total | pessoa", True),
        ("Teto do grupo", '=IF(B11="pessoa",B10*B8,B10)', FMT_EUR0, "fórmula", False),
        ("Reserva para imprevistos (%)", (o.get("reserva") or {}).get("percentual", 10) / 100, FMT_PCT, "sobre o total escolhido", True),
        ("Margem do cartão fora do euro", float(o.get("margemCartao") or 0), FMT_PCT, "0% num Revolut/Wise; 1–3% em banco tradicional", True),
        ("Opção padrão dos itens", "plano", None, "a coluna Opção da aba Itens decide item a item", False),
        ("Pesquisa de", iso(v.get("atualizadoEm")), FMT_DATA, None, False),
        ("Status", v.get("status", "planejando"), None, None, False),
    ]
    for i, (rot, val, fmt, nota, chave) in enumerate(linhas, start=3):
        rotulo(ws.cell(row=i, column=1), rot)
        c = ws.cell(row=i, column=2)
        if isinstance(val, str) and val.startswith("="):
            formula(c, val, fmt)
        else:
            entrada(c, val, fmt, chave)
        if nota:
            rotulo(ws.cell(row=i, column=5), nota, cor=CINZA, italico=True)
    dv = DataValidation(type="list", formula1='"total,pessoa"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add("B11")

    r = 20
    rotulo(ws.cell(row=r, column=1), "Câmbio (EUR por 1 unidade da moeda)", negrito=True)
    cabecalho(ws, r + 1, ["Moeda", "EUR por 1", "Data", "", "Fonte"])
    moedas = {base: (1.0, iso(v.get("atualizadoEm")), "moeda de casa")}
    for c in o.get("cambioReferencia", []):
        moedas[c["de"]] = (float(c["taxa"]), iso(c.get("em")), "cotação comercial na data")
    for it in extrair_itens(v):
        moedas.setdefault(it["moeda"], (None, None, "SEM COTAÇÃO — preencha"))
    for d in v.get("destinos", []):
        if d.get("moeda"):
            moedas.setdefault(d["moeda"], (None, None, "SEM COTAÇÃO — preencha"))
    r0 = r + 2
    for i, (m, (taxa, em, fnt)) in enumerate(moedas.items()):
        rr = r0 + i
        entrada(ws.cell(row=rr, column=1), m)
        entrada(ws.cell(row=rr, column=2), taxa, FMT_TAXA, chave=(m != base))
        entrada(ws.cell(row=rr, column=3), em, FMT_DATA)
        rotulo(ws.cell(row=rr, column=5), fnt, cor=CINZA, italico=True)
    abas.cambio = (r0, r0 + len(moedas) - 1)

    r = abas.cambio[1] + 3
    rotulo(ws.cell(row=r, column=1), "Categorias", negrito=True)
    rotulo(ws.cell(row=r, column=5), "Fator econômico: preço que sobra na opção econômica quando o item não tem uma explícita (0,8 = corta 20%). Intocável força 1.", cor=CINZA, italico=True)
    cabecalho(ws, r + 1, ["ID", "Nome", "Fator econômico", "Intocável", "De onde veio o número"])
    _, cats = itens_com_restante(v, pessoas)
    r0 = r + 2
    for i, c in enumerate(cats):
        rr = r0 + i
        entrada(ws.cell(row=rr, column=1), c["id"])
        entrada(ws.cell(row=rr, column=2), c["nome"])
        entrada(ws.cell(row=rr, column=3), c["fator"], "0.00")
        entrada(ws.cell(row=rr, column=4), "sim" if c["intocavel"] else "não")
        rotulo(ws.cell(row=rr, column=5), c["observacao"], cor=CINZA)
    abas.cats = (r0, r0 + len(cats) - 1)
    abas.cat_ids = [c["id"] for c in cats]
    dv2 = DataValidation(type="list", formula1='"sim,não"', allow_blank=False)
    ws.add_data_validation(dv2)
    dv2.add(f"D{r0}:D{abas.cats[1]}")
    ws.freeze_panes = "A3"


def aba_itens(ws, v: dict, abas: Abas, pessoas: int) -> list[dict]:
    itens, cats = itens_com_restante(v, pessoas)
    P = abas.P
    ci, cf = abas.cambio
    ki, kf = abas.cats
    base_ref = ref_param(abas, "B9")
    pessoas_ref = ref_param(abas, "B8")
    margem_ref = ref_param(abas, "B14")
    moedas = f"{P}!$A${ci}:$A${cf}"
    taxas = f"{P}!$B${ci}:$B${cf}"
    cat_ids = f"{P}!$A${ki}:$A${kf}"
    cat_fator = f"{P}!$C${ki}:$C${kf}"
    cat_intoc = f"{P}!$D${ki}:$D${kf}"
    C = COL

    cabecalho(ws, 1, [
        "ID", "Categoria", "Item", "Quando", "Qtd", "Unidade", "Moeda", "Opção ▼",
        "Preço econômico", "Preço plano", "Preço upgrade", "Preço escolhido",
        "× (pessoas/noites)", "Câmbio + margem",
        "Total econômico (EUR)", "Total plano (EUR)", "Total upgrade (EUR)", "Total escolhido (EUR)",
        "Intocável", "Dia", "Reservar / comprar", "O que é o econômico", "O que é o upgrade", "Fonte / observação",
    ], [18, 15, 44, 11, 5, 12, 7, 11, 13, 12, 13, 13, 10, 11, 15, 15, 15, 16, 9, 5, 34, 40, 40, 60])

    dv_un = DataValidation(type="list", formula1='"pessoa,total,noite,pessoa-noite"', allow_blank=False)
    dv_cat = DataValidation(type="list", formula1=f"={cat_ids}", allow_blank=False)
    dv_op = DataValidation(type="list", formula1='"econômico,plano,upgrade"', allow_blank=False)
    for dv in (dv_un, dv_cat, dv_op):
        ws.add_data_validation(dv)

    ordem = {c["id"]: i for i, c in enumerate(cats)}
    itens.sort(key=lambda it: (ordem.get(it["categoria"], 99), it.get("quando") or date.max, it["id"]))
    base = v.get("moedaBase") or "EUR"

    r = 2
    for it in itens:
        fmt_preco = FMT_EUR if it["moeda"] == base else FMT_NUM
        entrada(ws[f"{C['id']}{r}"], it["id"])
        entrada(ws[f"{C['cat']}{r}"], it["categoria"])
        entrada(ws[f"{C['item']}{r}"], it["item"])
        entrada(ws[f"{C['quando']}{r}"], it.get("quando"), FMT_DATA)
        entrada(ws[f"{C['qtd']}{r}"], it["qtd"], "0")
        entrada(ws[f"{C['unid']}{r}"], it["unidade"])
        entrada(ws[f"{C['moeda']}{r}"], it["moeda"])
        entrada(ws[f"{C['opcao']}{r}"], it.get("opcao", "plano"))
        ws[f"{C['opcao']}{r}"].fill = ESCOLHA
        # preço econômico: explícito, ou plano × fator da categoria
        if it.get("economico") is not None:
            entrada(ws[f"{C['p_eco']}{r}"], it["economico"], fmt_preco)
        else:
            formula(ws[f"{C['p_eco']}{r}"], f"={C['p_plano']}{r}*IFERROR(INDEX({cat_fator},MATCH({C['cat']}{r},{cat_ids},0)),1)", fmt_preco)
        entrada(ws[f"{C['p_plano']}{r}"], it["plano"], fmt_preco)
        # upgrade: explícito, ou igual ao plano (sem upgrade conhecido)
        if it.get("upgrade") is not None:
            entrada(ws[f"{C['p_up']}{r}"], it["upgrade"], fmt_preco)
        else:
            formula(ws[f"{C['p_up']}{r}"], f"={C['p_plano']}{r}", fmt_preco)
        formula(ws[f"{C['p_esc']}{r}"],
                f'=IF({C["opcao"]}{r}="econômico",{C["p_eco"]}{r},IF({C["opcao"]}{r}="upgrade",{C["p_up"]}{r},{C["p_plano"]}{r}))',
                fmt_preco, negrito=True)
        formula(ws[f"{C['mult']}{r}"], f'={C["qtd"]}{r}*IF(OR({C["unid"]}{r}="pessoa",{C["unid"]}{r}="pessoa-noite"),{pessoas_ref},1)', "0", VERDE)
        formula(ws[f"{C['cambio']}{r}"],
                f"=IFERROR(INDEX({taxas},MATCH({C['moeda']}{r},{moedas},0)),1)*(1+IF({C['moeda']}{r}={base_ref},0,{margem_ref}))",
                FMT_TAXA, VERDE)
        for tcol, pcol in ((C["t_eco"], C["p_eco"]), (C["t_plano"], C["p_plano"]), (C["t_up"], C["p_up"])):
            formula(ws[f"{tcol}{r}"], f"={C['mult']}{r}*{pcol}{r}*{C['cambio']}{r}", FMT_EUR)
        formula(ws[f"{C['t_esc']}{r}"], f"={C['mult']}{r}*{C['p_esc']}{r}*{C['cambio']}{r}", FMT_EUR, negrito=True)
        formula(ws[f"{C['intoc']}{r}"], f'=IFERROR(INDEX({cat_intoc},MATCH({C["cat"]}{r},{cat_ids},0)),"não")', None, VERDE)
        entrada(ws[f"{C['dia']}{r}"], it.get("dia"), "0")
        link(ws[f"{C['link']}{r}"], it.get("link"), _rotulo_link(it.get("link")))
        rotulo(ws[f"{C['d_eco']}{r}"], it.get("d_eco") or ("— (preço plano × fator da categoria)" if it.get("economico") is None and not it.get("estimativa") else ""), cor=CINZA)
        rotulo(ws[f"{C['d_up']}{r}"], it.get("d_up") or ("— (sem upgrade pesquisado)" if it.get("upgrade") is None and not it.get("estimativa") else ""), cor=CINZA)
        rotulo(ws[f"{C['fonte']}{r}"], it.get("fonte", ""), cor=CINZA)
        if it.get("estimativa"):
            for col in "ABCDEFGHIJK":
                cel = ws[f"{col}{r}"]
                cel.font = Font(name=FONTE, color=cel.font.color, italic=True, size=10)
            ws[f"{C['p_plano']}{r}"].comment = Comment(
                "Estimativa da categoria que não está detalhada em itens (vem de orcamento.categorias.previsto no trip.json). "
                "Ao detalhar um item novo, reduza este valor para não contar duas vezes.", "planilha-viagem")
        for col in (C["opcao"], C["unid"], C["cat"]):
            pass
        abas.linhas_itens[it["id"]] = r
        r += 1
    abas.fim_itens = r - 1
    dv_un.add(f"{C['unid']}2:{C['unid']}{r + 200}")
    dv_cat.add(f"{C['cat']}2:{C['cat']}{r + 200}")
    dv_op.add(f"{C['opcao']}2:{C['opcao']}{r + 200}")
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = f"A1:{C['fonte']}{r - 1}"

    r += 1
    rotulo(ws[f"C{r}"], "Como usar", negrito=True)
    for t in [
        "Opção ▼ (coluna H, fundo creme) escolhe o que entra no total: econômico, plano ou upgrade. É o único lugar que você precisa mexer para montar o seu cenário.",
        "As três colunas de preço são por unidade (por pessoa, por noite ou total, conforme a Unidade). Azul = pesquisado; preto = derivado do plano.",
        "Preço econômico sem número explícito = preço do plano × fator da categoria (aba Parâmetros). Preço upgrade sem número = igual ao plano.",
        "Reservar / comprar (coluna U) abre o site oficial ou o comparador. Compre pelo link; a taxa de serviço de agregador entra no último passo.",
        "Linhas em itálico são estimativas da categoria: ao detalhar um item novo, reduza a estimativa para não contar duas vezes.",
        "Para acrescentar um item: copie uma linha inteira (as fórmulas de L a S vêm junto) e troque ID, categoria, preços e link.",
    ]:
        r += 1
        rotulo(ws[f"C{r}"], t, cor=CINZA)
    return itens


def _rotulo_link(url: str | None) -> str | None:
    if not url:
        return None
    dom = url.split("//")[-1].split("/")[0]
    return dom.removeprefix("www.")


def _sumifs_cat(abas: Abas, col: str, cat_cel: str) -> str:
    I = abas.I
    return f"=SUMIFS({I}!${col}:${col},{I}!$B:$B,{cat_cel})"


def aba_resumo(ws, v: dict, abas: Abas, cats: list[dict], com_real: bool = True) -> dict:
    P, I, R = abas.P, abas.I, abas.R
    C = COL
    titulo(ws, v.get("titulo", v["slug"]), v.get("subtitulo") or "")
    ws.column_dimensions["A"].width = 30
    for col in "BCDEFG":
        ws.column_dimensions[col].width = 17
    ws.column_dimensions["H"].width = 10

    rotulo(ws["A4"], "Período")
    formula(ws["B4"], f"={P}!B4", FMT_DATA, VERDE)
    formula(ws["C4"], f"={P}!B5", FMT_DATA, VERDE)
    rotulo(ws["A5"], "Pessoas · noites · dias")
    formula(ws["B5"], f"={P}!B8", "0", VERDE)
    formula(ws["C5"], f"={P}!B6", "0", VERDE)
    formula(ws["D5"], f"={P}!B7", "0", VERDE)

    h = 7
    cabecalho(ws, h, ["Categoria", "Tudo econômico", "Plano", "Escolhido ★", "Tudo upgrade", "Real", "Real − Escolhido", "Intocável"])
    ki, _ = abas.cats
    r = h + 1
    for i, c in enumerate(cats):
        pr = ki + i
        formula(ws.cell(row=r, column=1), f"={P}!B{pr}", None, VERDE)
        formula(ws.cell(row=r, column=2), _sumifs_cat(abas, C["t_eco"], f"{P}!$A${pr}"), FMT_EUR, VERDE)
        formula(ws.cell(row=r, column=3), _sumifs_cat(abas, C["t_plano"], f"{P}!$A${pr}"), FMT_EUR, VERDE)
        formula(ws.cell(row=r, column=4), _sumifs_cat(abas, C["t_esc"], f"{P}!$A${pr}"), FMT_EUR, VERDE, negrito=True)
        formula(ws.cell(row=r, column=5), _sumifs_cat(abas, C["t_up"], f"{P}!$A${pr}"), FMT_EUR, VERDE)
        if com_real:
            formula(ws.cell(row=r, column=6), f"=SUMIFS({R}!$G:$G,{R}!$B:$B,{P}!$A${pr})", FMT_EUR, VERDE)
            formula(ws.cell(row=r, column=7), f'=IF(F{r}=0,"",F{r}-D{r})', FMT_EUR)
        formula(ws.cell(row=r, column=8), f"={P}!D{pr}", None, VERDE)
        if c["intocavel"]:
            for col in range(1, 9):
                ws.cell(row=r, column=col).fill = INTOCAVEL
        r += 1
    r_ini, r_fim = h + 1, r - 1
    linhas = {}

    def linha(nome, fs, fmt=FMT_EUR, negrito=False, fill=None):
        nonlocal r
        rotulo(ws.cell(row=r, column=1), nome, negrito=negrito)
        for col, f in zip("BCDEF", fs):
            if f:
                formula(ws[f"{col}{r}"], f, fmt, negrito=negrito or col == "D")
        if fill:
            for col in range(1, 9):
                ws.cell(row=r, column=col).fill = fill
        linhas[nome] = r
        r += 1
        return r - 1

    cols = "BCDE" + ("F" if com_real else "")
    t = linha("Total previsto", [f"=SUM({c}{r_ini}:{c}{r_fim})" for c in cols], negrito=True, fill=SUBTOTAL)
    linha("Reserva para imprevistos", [f"={c}{t}*{P}!$B$13" for c in "BCDE"])
    tr = linha("Total com reserva", [f"={c}{t}+{c}{t + 1}" for c in "BCDE"], negrito=True, fill=SUBTOTAL)
    r += 1
    linha("Por pessoa", [f"={c}{t}/{P}!$B$8" for c in cols])
    linha("Por dia (grupo)", [f"={c}{t}/{P}!$B$7" for c in cols])
    linha("Por pessoa por dia", [f"={c}{t}/{P}!$B$8/{P}!$B$7" for c in cols])
    tcols = {"B": C["t_eco"], "C": C["t_plano"], "D": C["t_esc"], "E": C["t_up"]}
    gs = linha("Gasto em solo (sem voos)", [f'={c}{t}-SUMIFS({I}!${tcols[c]}:${tcols[c]},{I}!$B:$B,"voos")' for c in "BCDE"])
    linha("Alimentação / gasto em solo",
          [f'=IF({c}{gs}=0,0,SUMIFS({I}!${tcols[c]}:${tcols[c]},{I}!$B:$B,"alimentacao")/{c}{gs})' for c in "BCDE"], fmt=FMT_PCT)
    rotulo(ws.cell(row=r - 1, column=8), "meta 15–20%", cor=CINZA, italico=True)
    r += 1
    teto = linha("Teto do grupo", [f"={P}!$B$12"] * 4, negrito=True)
    folga = linha("Folga contra o teto", [f"={c}{teto}-{c}{tr}" for c in "BCDE"], negrito=True)
    linha("Uso do teto", [f"=IF({c}{teto}=0,0,{c}{tr}/{c}{teto})" for c in "BCDE"], fmt=FMT_PCT)
    rotulo(ws.cell(row=r, column=1), "Veredito", negrito=True)
    for col in "BCDE":
        formula(ws[f"{col}{r}"], f'=IF({col}{folga}>=0,"cabe, sobram "&TEXT({col}{folga},"#,##0")&" €","faltam "&TEXT(-{col}{folga},"#,##0")&" €")', None, negrito=True)
    r += 1
    rotulo(ws.cell(row=r, column=1), "Itens em upgrade / econômico")
    formula(ws[f"D{r}"], f'=COUNTIF({I}!${C["opcao"]}:${C["opcao"]},"upgrade")&" up · "&COUNTIF({I}!${C["opcao"]}:${C["opcao"]},"econômico")&" eco"', None, VERDE)
    r += 2

    rotulo(ws.cell(row=r, column=1), "Como ler", negrito=True)
    for tx in [
        "Escolhido ★ é o seu cenário: soma a opção marcada em cada linha da aba Itens (coluna Opção ▼). Tudo econômico e Tudo upgrade são os dois extremos.",
        "Plano = o roteiro como foi pesquisado. Real soma o que você lançar na aba Real durante a viagem; Real − Escolhido mostra onde estourou.",
        "Linhas rosa são intocáveis: a razão de ser da viagem. Cortar ali é planejar outra viagem.",
        "Nada aqui é digitado: tudo vem de Itens, Parâmetros e Real. Para mudar um número, mude lá.",
    ]:
        r += 1
        rotulo(ws.cell(row=r, column=1), tx, cor=CINZA)

    ch = BarChart()
    ch.type = "col"
    ch.title = "Por categoria: econômico × plano × escolhido × upgrade"
    ch.y_axis.numFmt = "#,##0"
    ch.add_data(Reference(ws, min_col=2, min_row=h, max_col=5, max_row=r_fim), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=h + 1, max_row=r_fim))
    ch.height, ch.width = 8, 24
    ws.add_chart(ch, "J4")
    ws.freeze_panes = "A8"
    return dict(total=t, total_reserva=tr, folga=folga, teto=teto)


def aba_roteiro(ws, v: dict, abas: Abas):
    I, C = abas.I, COL
    titulo(ws, "Roteiro dia a dia", "Custo de cada bloco vem da aba Itens (opção escolhida). Marque Feito durante a viagem.")
    cabecalho(ws, 4, ["Dia", "Data", "Semana", "Hora", "Bloco", "Tipo", "Custo escolhido (EUR)", "Opção", "Reservar antes", "Feito", "Link", "Notas / plano B"],
              [5, 12, 10, 7, 50, 13, 16, 11, 12, 7, 28, 70])
    dv = DataValidation(type="list", formula1='"sim,"', allow_blank=True)
    ws.add_data_validation(dv)
    r = 5
    for i, dia in enumerate(v.get("dias", []), start=1):
        d = iso(dia.get("data"))
        rotulo(ws.cell(row=r, column=1), i, negrito=True)
        entrada(ws.cell(row=r, column=2), d, FMT_DATA)
        rotulo(ws.cell(row=r, column=3), SEMANA[d.weekday()] if d else "")
        rotulo(ws.cell(row=r, column=5), (dia.get("titulo") or "") + ("   [cortável]" if dia.get("cortavel") else ""), negrito=True)
        rotulo(ws.cell(row=r, column=12), dia.get("notas") or "", cor=CINZA)
        for col in range(1, 13):
            ws.cell(row=r, column=col).fill = SUBTOTAL
        r0 = r + 1
        r = r0
        for b in dia.get("blocos", []):
            rotulo(ws.cell(row=r, column=4), b.get("hora") or "")
            rotulo(ws.cell(row=r, column=5), b.get("titulo") or "")
            rotulo(ws.cell(row=r, column=6), b.get("tipo") or "")
            bid = f"dia{i}:{b.get('id', '')}"
            if bid in abas.linhas_itens:
                li = abas.linhas_itens[bid]
                formula(ws.cell(row=r, column=7), f"={I}!{C['t_esc']}{li}", FMT_EUR, VERDE)
                formula(ws.cell(row=r, column=8), f"={I}!{C['opcao']}{li}", None, VERDE)
            rotulo(ws.cell(row=r, column=9), "sim" if b.get("reservaNecessaria") else "")
            dv.add(ws.cell(row=r, column=10).coordinate)
            link(ws.cell(row=r, column=11), b.get("link"), _rotulo_link(b.get("link")))
            nota = (b.get("notas") or "")
            if b.get("alternativa"):
                nota += ("  ·  " if nota else "") + "Plano B: " + b["alternativa"]
            rotulo(ws.cell(row=r, column=12), nota, cor=CINZA)
            r += 1
        rotulo(ws.cell(row=r, column=5), f"Custo do dia {i}", italico=True)
        formula(ws.cell(row=r, column=7), f"=SUM(G{r0}:G{r - 1})" if r > r0 else "=0", FMT_EUR, negrito=True)
        r += 2
    rotulo(ws.cell(row=r, column=5), "Total dos blocos com custo", negrito=True)
    formula(ws.cell(row=r, column=7), f'=SUMIFS({I}!${C["t_esc"]}:${C["t_esc"]},{I}!${C["dia"]}:${C["dia"]},">0")', FMT_EUR, VERDE, negrito=True)
    ws.freeze_panes = "A5"


URGENCIA = {"imediato": "Agora", "data-exata": "Hora marcada", "sorteio": "Sorteio", "um-mes": "Um mês antes",
            "duas-semanas": "2 a 4 semanas", "ultima-semana": "Última semana"}


def aba_reservas(ws, v: dict):
    titulo(ws, "Reservar antes", "Em ordem de prazo. Marque Status quando fizer e anote o valor pago para o Real.")
    cabecalho(ws, 4, ["O quê", "Urgência", "Prazo", "Dias restantes", "Status", "Valor pago (EUR)", "Reservar / comprar", "Observação"],
              [52, 14, 12, 12, 12, 15, 34, 90])
    dv = DataValidation(type="list", formula1='"pendente,feita,não precisa"', allow_blank=False)
    ws.add_data_validation(dv)
    res = sorted(v.get("reservas", []), key=lambda x: x.get("prazo") or "9999")
    r = 5
    for x in res:
        rotulo(ws.cell(row=r, column=1), x.get("oQue", ""))
        rotulo(ws.cell(row=r, column=2), URGENCIA.get(x.get("urgencia"), x.get("urgencia") or ""))
        entrada(ws.cell(row=r, column=3), iso(x.get("prazo")), FMT_DATA)
        formula(ws.cell(row=r, column=4), f'=IF(C{r}="","",C{r}-TODAY())', "0")
        entrada(ws.cell(row=r, column=5), "pendente")
        dv.add(f"E{r}")
        entrada(ws.cell(row=r, column=6), None, FMT_EUR)
        link(ws.cell(row=r, column=7), x.get("onde"), _rotulo_link(x.get("onde")))
        obs = x.get("observacao") or ""
        av = (x.get("abreVendaEm") or {}).get("emCasa")
        if av:
            obs = f"Abre à venda {av[:10]} às {av[11:16]} (casa).  " + obs
        if x.get("nominativo"):
            obs = f"Nominativo: {x['nominativo']}.  " + obs
        rotulo(ws.cell(row=r, column=8), obs, cor=CINZA)
        r += 1
    rotulo(ws["A2"], "Progresso", negrito=True)
    formula(ws["B2"], f'=COUNTIF(E5:E{r - 1},"feita")&" de "&COUNTA(A5:A{r - 1})&" feitas"', None, negrito=True)
    ws.freeze_panes = "A5"


def aba_real(ws, v: dict, abas: Abas):
    P = abas.P
    ci, cf = abas.cambio
    ki, kf = abas.cats
    titulo(ws, "Gastos reais", "Lance cada gasto durante a viagem. A coluna EUR converte pela aba Parâmetros e o Resumo compara com o escolhido.")
    cabecalho(ws, 4, ["Data", "Categoria", "Descrição", "Valor", "Moeda", "Ignorar", "EUR", "Pago por"],
              [12, 16, 48, 12, 7, 8, 14, 14])
    dv_cat = DataValidation(type="list", formula1=f"={P}!$A${ki}:$A${kf}", allow_blank=True)
    dv_ig = DataValidation(type="list", formula1='"sim,"', allow_blank=True)
    ws.add_data_validation(dv_cat)
    ws.add_data_validation(dv_ig)
    base = v.get("moedaBase") or "EUR"
    entrada(ws["A5"], iso(v["periodo"]["inicio"]), FMT_DATA)
    entrada(ws["B5"], "alimentacao")
    entrada(ws["C5"], "(exemplo — apague ou sobrescreva) Jantar no bouillon")
    entrada(ws["D5"], 38.5, FMT_NUM)
    entrada(ws["E5"], base)
    entrada(ws["F5"], "sim")
    entrada(ws["H5"], "cartão")
    for r in range(5, 206):
        formula(ws.cell(row=r, column=7),
                f'=IF(OR(D{r}="",F{r}="sim"),0,D{r}*IFERROR(INDEX({P}!$B${ci}:$B${cf},MATCH(E{r},{P}!$A${ci}:$A${cf},0)),1)'
                f'*(1+IF(E{r}={P}!$B$9,0,{P}!$B$14)))', FMT_EUR, VERDE)
        dv_cat.add(f"B{r}")
        dv_ig.add(f"F{r}")
    rotulo(ws["A2"], "Total real", negrito=True)
    formula(ws["B2"], "=SUM(G5:G205)", FMT_EUR, negrito=True)
    ws.freeze_panes = "A5"


def aba_cortes(ws, v: dict, abas: Abas, itens: list[dict], resumo_linhas: dict):
    I, RS, C = abas.I, q(abas.resumo), COL
    titulo(ws, "Onde cortar, na ordem que menos dói",
           "Economia de trocar cada item do escolhido para o econômico, do maior corte ao menor. Intocáveis ficam por último e em rosa.")
    cabecalho(ws, 4, ["Item", "Categoria", "Opção atual", "Escolhido (EUR)", "Econômico (EUR)", "Economia", "Acumulado", "Intocável", "Já basta?"],
              [50, 15, 11, 15, 15, 12, 12, 10, 10])
    rotulo(ws["A2"], "Falta para caber no teto (escolhido com reserva):")
    formula(ws["D2"], f"=MAX(0,-{RS}!D{resumo_linhas['folga']})", FMT_EUR, VERDE, negrito=True)
    pessoas = pessoas_de(v)
    cats = {c["id"]: c for c in itens_com_restante(v, pessoas)[1]}

    def corte(it):
        c = cats.get(it["categoria"], {})
        if c.get("intocavel"):
            return -1
        eco = it["economico"] if it.get("economico") is not None else it["plano"] * c.get("fator", 1)
        mult = it["qtd"] * (pessoas if it["unidade"] in ("pessoa", "pessoa-noite") else 1)
        return (it["plano"] - eco) * mult

    r = 5
    for it in sorted(itens, key=corte, reverse=True):
        li = abas.linhas_itens[it["id"]]
        formula(ws.cell(row=r, column=1), f"={I}!{C['item']}{li}", None, VERDE)
        formula(ws.cell(row=r, column=2), f"={I}!{C['cat']}{li}", None, VERDE)
        formula(ws.cell(row=r, column=3), f"={I}!{C['opcao']}{li}", None, VERDE)
        formula(ws.cell(row=r, column=4), f"={I}!{C['t_esc']}{li}", FMT_EUR, VERDE)
        formula(ws.cell(row=r, column=5), f"={I}!{C['t_eco']}{li}", FMT_EUR, VERDE)
        formula(ws.cell(row=r, column=6), f"=MAX(0,D{r}-E{r})", FMT_EUR)
        formula(ws.cell(row=r, column=7), f"=SUM($F$5:F{r})", FMT_EUR)
        formula(ws.cell(row=r, column=8), f"={I}!{C['intoc']}{li}", None, VERDE)
        prev = f"G{r - 1}<$D$2" if r > 5 else "TRUE"
        formula(ws.cell(row=r, column=9), f'=IF(AND($D$2>0,G{r}>=$D$2,{prev}),"← aqui","")', None, negrito=True)
        if cats.get(it["categoria"], {}).get("intocavel"):
            for col in range(1, 10):
                ws.cell(row=r, column=col).fill = INTOCAVEL
        r += 1
    r += 1
    rotulo(ws.cell(row=r, column=1), "A ordem foi calculada na geração; os valores são vivos. Se você mudar preços na aba Itens, a ordem pode ficar desatualizada — a economia por linha não.", cor=CINZA, italico=True)
    ws.freeze_panes = "A5"


def montar_viagem(wb, v: dict, abas: Abas | None = None, com_extras: bool = True):
    abas = abas or Abas()
    pessoas = pessoas_de(v)
    ws_res = wb.create_sheet(abas.resumo)
    ws_par = wb.create_sheet(abas.params)
    ws_it = wb.create_sheet(abas.itens)
    aba_parametros(ws_par, v, abas, pessoas)
    itens = aba_itens(ws_it, v, abas, pessoas)
    _, cats = itens_com_restante(v, pessoas)
    if com_extras:
        ws_rot = wb.create_sheet("Roteiro")
        ws_resv = wb.create_sheet("Reservas")
        ws_real = wb.create_sheet(abas.real)
        ws_cort = wb.create_sheet("Cortes")
        linhas = aba_resumo(ws_res, v, abas, cats, com_real=True)
        aba_roteiro(ws_rot, v, abas)
        aba_reservas(ws_resv, v)
        aba_real(ws_real, v, abas)
        aba_cortes(ws_cort, v, abas, itens, linhas)
    else:
        linhas = aba_resumo(ws_res, v, abas, cats, com_real=False)
    return abas, itens, cats, linhas
